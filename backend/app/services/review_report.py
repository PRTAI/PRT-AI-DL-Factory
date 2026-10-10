"""样例审核阶段 ④：把逐项结论汇总成 report.json，再渲染成固定格式的 docx。

报告格式是**唯一交付形态**，不允许按次生成不同的排版：所有样式常量集中在本文件
顶部，``build_report()`` 只管内容、``render_docx()`` 只管排版，两者用一个 dict 连接
（report.json 先落盘，docx 是它的渲染）。这样报告长什么样可以被测试断言。

骨架、配色与列宽量自 ``返修报告_铂瑞缇_金融_0923.pdf``。
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..models import ReviewCheck, ReviewFinding, ReviewTask, ReviewVerdict

# ----------------------------------------------------------------- 固定样式 ---
FONT_CN = "微软雅黑"
FONT_EN = "Arial"
COLOR_HEADER_BG = "34495E"   # 表头底色
COLOR_HEADER_FG = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ZEBRA = "F8F9FA"       # 表体隔行底色
COLOR_GRID = "BDC3C7"        # 网格线
COLOR_TEXT = RGBColor(0x2C, 0x3E, 0x50)
COLOR_MUTED = RGBColor(0x7F, 0x8C, 0x8D)
# 三列宽度（pt），量自参考报告的竖网格线 x=51/110/297/544
COL_WIDTHS_PT = (59, 187, 247)
COL_WIDTHS_4_PT = (59, 187, 128, 119)  # 受检清单多一列，后两列从第三列匀出来

VERDICT_LABEL = {
    ReviewVerdict.PASS: "通过",
    ReviewVerdict.FAIL: "不合格",
    ReviewVerdict.NEEDS_REVIEW: "需人工复核",
}

HEAD_DETAIL = ("题号", "源数据相对位置", "具体问题")
HEAD_TASKS = ("题号", "源数据相对位置", "结论", "不合格项")
HEAD_MANUAL = ("题号", "源数据相对位置", "检查项", "待确认点")

SEC_NOTE = "一、说明"
SEC_DETAIL = "二、不合格明细"
SEC_TASKS = "三、受检题目清单"
SEC_MANUAL = "四、需人工复核"


# --------------------------------------------------------------------------- #
# report.json
# --------------------------------------------------------------------------- #
def build_report(
    task: ReviewTask,
    manifest: dict[str, Any],
    checks: list[ReviewCheck],
    findings: list[ReviewFinding],
) -> dict[str, Any]:
    """逐题汇总。任一项不合格 → 不合格；否则有待确认 → 需人工复核；都没有 → 通过。"""
    tasks = manifest.get("tasks") or []
    order = {t["task_id"]: i for i, t in enumerate(tasks)}
    rel_of = {t["task_id"]: t.get("rel_path", "") for t in tasks}
    check_of = {c.code: c for c in checks}

    fails: dict[str, list[ReviewFinding]] = {}
    manual: dict[str, list[ReviewFinding]] = {}
    for finding in findings:
        if finding.verdict is ReviewVerdict.FAIL:
            fails.setdefault(finding.task_id, []).append(finding)
        elif finding.verdict is ReviewVerdict.NEEDS_REVIEW:
            manual.setdefault(finding.task_id, []).append(finding)

    task_rows: list[dict[str, Any]] = []
    for item in tasks:
        tid = item["task_id"]
        if tid in fails:
            verdict = ReviewVerdict.FAIL
        elif tid in manual:
            verdict = ReviewVerdict.NEEDS_REVIEW
        else:
            verdict = ReviewVerdict.PASS
        task_rows.append(
            {
                "task_id": tid,
                "rel_path": item.get("rel_path", ""),
                "verdict": verdict.value,
                "verdict_label": VERDICT_LABEL[verdict],
                "fail_codes": sorted({f.check_code for f in fails.get(tid, [])}),
            }
        )

    n_tasks = len(task_rows)
    n_fail = sum(1 for r in task_rows if r["verdict"] == ReviewVerdict.FAIL.value)
    n_manual = sum(1 for r in task_rows if r["verdict"] == ReviewVerdict.NEEDS_REVIEW.value)
    n_pass = n_tasks - n_fail - n_manual

    # 二、不合格明细：按检查项分节，只留有不合格的项
    sections: list[dict[str, Any]] = []
    by_check: dict[str, list[ReviewFinding]] = {}
    for finding in findings:
        if finding.verdict is ReviewVerdict.FAIL:
            by_check.setdefault(finding.check_code, []).append(finding)
    for code in sorted(by_check, key=lambda c: (check_of[c].code if c in check_of else c)):
        check = check_of.get(code)
        rows = sorted(by_check[code], key=lambda f: order.get(f.task_id, 1 << 30))
        sections.append(
            {
                "code": code,
                "title": check.title if check else code,
                "source": check.source if check else "",
                "requirement": check.requirement if check else "",
                "rows": [
                    {
                        "task_id": f.task_id,
                        "rel_path": f.rel_path or rel_of.get(f.task_id, ""),
                        "problems": list(f.problems),
                    }
                    for f in rows
                ],
            }
        )

    manual_rows = [
        {
            "task_id": f.task_id,
            "rel_path": f.rel_path or rel_of.get(f.task_id, ""),
            "check": f"{f.check_code} {check_of[f.check_code].title}" if f.check_code in check_of else f.check_code,
            "problems": list(f.problems),
        }
        for f in sorted(
            (f for f in findings if f.verdict is ReviewVerdict.NEEDS_REVIEW),
            key=lambda f: (order.get(f.task_id, 1 << 30), f.check_code),
        )
    ]

    return {
        "batch": task.name,
        "supplier": task.supplier,
        "spec_title": task.spec_title,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "model_label": task.model_label,
        "stats": {
            "n_tasks": n_tasks,
            "n_pass": n_pass,
            "n_fail": n_fail,
            "n_needs_review": n_manual,
            "n_checks": len(checks),
            "pass_rate": (n_pass / n_tasks) if n_tasks else 0.0,
        },
        "checks": [
            {"code": c.code, "title": c.title, "source": c.source, "requirement": c.requirement}
            for c in checks
        ],
        "sections": sections,
        "tasks": task_rows,
        "needs_review": manual_rows,
        "notes": list(task.notes),
    }


def stats_line(stats: dict[str, Any]) -> str:
    return (
        f"受检 {stats['n_tasks']} 题；通过 {stats['n_pass']} 题；"
        f"不合格 {stats['n_fail']} 题；需人工复核 {stats['n_needs_review']} 题；"
        f"通过率 {stats['pass_rate'] * 100:.1f}%。"
    )


def write_report_json(report: dict[str, Any], path: Path) -> None:
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


# --------------------------------------------------------------------------- #
# docx
# --------------------------------------------------------------------------- #
def render_docx(report: dict[str, Any], path: Path) -> None:
    doc = Document()
    _setup(doc)

    _title(doc, f"{report['batch']} 静态验收不合格返修报告")
    _subtitle(doc, f"供应商：{report['supplier']}" if report.get("supplier") else "")

    _heading(doc, SEC_NOTE)
    spec = report.get("spec_title") or "评测题包交付规范"
    _body(doc, f"依据：《{spec}》。同一题可出现在多个验收项中，题数不可相加。")
    _body(doc, stats_line(report["stats"]))
    if report.get("notes"):
        _body(doc, "审核备注：" + "；".join(str(n) for n in report["notes"]), muted=True)

    _heading(doc, SEC_DETAIL)
    if not report["sections"]:
        _body(doc, "本次未发现不合格项。", muted=True)
    for section in report["sections"]:
        _check_heading(doc, section)
        _table(
            doc,
            HEAD_DETAIL,
            [
                [row["task_id"], row["rel_path"], "\n".join(row["problems"])]
                for row in section["rows"]
            ],
            COL_WIDTHS_PT,
        )

    _heading(doc, SEC_TASKS)
    _table(
        doc,
        HEAD_TASKS,
        [
            [row["task_id"], row["rel_path"], row["verdict_label"], "、".join(row["fail_codes"])]
            for row in report["tasks"]
        ],
        COL_WIDTHS_4_PT,
    )

    if report.get("needs_review"):
        _heading(doc, SEC_MANUAL)
        _body(doc, "以下题目证据不足，无法直接判定，请人工确认后再决定是否返修。", muted=True)
        _table(
            doc,
            HEAD_MANUAL,
            [
                [row["task_id"], row["rel_path"], row["check"], "\n".join(row["problems"])]
                for row in report["needs_review"]
            ],
            COL_WIDTHS_4_PT,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _setup(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Mm(20))
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(22)

    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = Pt(10)
    style.font.color.rgb = COLOR_TEXT
    _east_asian(style.element.get_or_add_rPr())
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25


def _east_asian(rPr) -> None:
    """中文字体得单独写 w:eastAsia，只设 w:ascii 的话 Word 会拿默认宋体去排中文。"""
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT_CN)


def _run(paragraph, text: str, *, size: float, bold: bool = False, color: RGBColor | None = None):
    run = paragraph.add_run(text)
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    _east_asian(run._element.get_or_add_rPr())
    return run


def _title(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(2)
    _run(para, text, size=16, bold=True)


def _subtitle(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(14)
    if text:
        _run(para, text, size=11, color=COLOR_MUTED)


def _heading(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(6)
    _run(para, text, size=13, bold=True)


def _body(doc: Document, text: str, *, muted: bool = False) -> None:
    para = doc.add_paragraph()
    _run(para, text, size=10, color=COLOR_MUTED if muted else None)


def _check_heading(doc: Document, section: dict[str, Any]) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(2)
    _run(para, f"{section['code']} {section['title']}", size=11, bold=True)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(4)
    pieces = []
    if section.get("source"):
        pieces.append(f"出处：{section['source']}")
    if section.get("requirement"):
        pieces.append(f"要求：{section['requirement']}")
    if pieces:
        _run(meta, "  ".join(pieces), size=9, color=COLOR_MUTED)


def _table(doc: Document, header: tuple[str, ...], rows: list[list[str]], widths: tuple[int, ...]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _grid(table)

    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        _shade(cell, COLOR_HEADER_BG)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        _run(para, text, size=9.5, bold=True, color=COLOR_HEADER_FG)

    for position, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if position % 2 == 1:
                _shade(cell, COLOR_ZEBRA)
            lines = str(value).split("\n") if value else [""]
            for line_no, line in enumerate(lines):
                para = cell.paragraphs[0] if line_no == 0 else cell.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                _run(para, line, size=9)

    for row in table.rows:  # 列宽必须逐格写，只设 table.columns 在 Word 里不生效
        for index, width in enumerate(widths[: len(header)]):
            row.cells[index].width = Pt(width)


def _shade(cell, hex_color: str) -> None:
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _grid(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.makeelement(qn(f"w:{edge}"), {})
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")  # 0.75pt
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), COLOR_GRID)
        borders.append(element)
    tbl_pr.append(borders)
